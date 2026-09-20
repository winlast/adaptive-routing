#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сверка чисел в статьях с сохранёнными данными.

Каждое число в тексте должно проверяться по файлу в data/. Проверка
существует потому, что однажды это уже было нарушено: таблица
блокировок собиралась по промежуточному прогону профилирования, а в
сохранённых данных лежал более поздний, и два значения оказались
переставлены местами. Ошибку нашёл читатель, а не автор.
"""
from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
DATA = BASE / "data"

problems: list[str] = []
checked = 0


def load(name):
    return json.loads((DATA / name).read_text(encoding="utf-8"))


def check(label: str, claim, actual, tol: float = 0.03) -> None:
    global checked
    checked += 1
    try:
        good = abs(float(claim) - float(actual)) <= max(
            tol * abs(float(actual)), 0.55)
    except (TypeError, ValueError):
        good = str(claim) == str(actual)
    if not good:
        problems.append(f"{label}: в статье {claim}, в данных {actual}")


def deck_text(path: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(path))
    return "\n".join(sh.text_frame.text for slide in prs.slides
                     for sh in slide.shapes if sh.has_text_frame)


def article_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    return "".join(re.findall(r"<w:t(?: [^>]*)?>(.*?)</w:t>", xml, re.S))


def main() -> int:
    detail = load("cost_grid_detail.json")

    def blk(ep, prm, worker, field):
        return next(r[field] for r in detail[ep][worker] if r["param"] == prm)

    # Таблица 1 статьи
    for ep, prm, vals in [
        ("/api/auth/verify", 1, (130, 9, 127, 125)),
        ("/api/report/generate", 45, (328, 76, 333, 222)),
        ("/api/report/generate", 75, (939, 268, 937, 638)),
        ("/api/feed", 400, (495, 9, 491, 10)),
    ]:
        for i, (w, f) in enumerate((("sync", "own_ms"), ("sync", "block_ms"),
                                    ("async", "own_ms"), ("async", "block_ms"))):
            check(f"табл.1 {ep}?{prm} {w}/{f}", vals[i], blk(ep, prm, w, f))

    # Точность оценок
    acc = load("estimator_accuracy.json")["overall"]
    for name, claim in [("среднее по маршруту", 80.3),
                        ("линейная поправка", 11.2),
                        ("степенной закон", 3.7), ("нейронная сеть", 3.4)]:
        check(f"точность: {name}", claim, acc[name])

    # Стоимость запроса и разброс внутри маршрута
    cost = load("endpoint_cost.json")

    def total(ep, prm):
        return next(r["total_ms"] for r in cost[ep] if r["param"] == prm)

    def cpu(ep, prm):
        return next(r["cpu_ms"] for r in cost[ep] if r["param"] == prm)

    check("search limit=10 полное", 10, total("/api/search", 10))
    check("search limit=10 вычисления", 4, cpu("/api/search", 10))
    check("search limit=600 полное", 640, total("/api/search", 600))
    check("search limit=600 вычисления", 433, cpu("/api/search", 600))
    for ep, claim in (("/api/search", 61), ("/api/report/generate", 85)):
        rows = [r["total_ms"] for r in cost[ep]]
        check(f"разброс внутри {ep}", claim, max(rows) / min(rows), tol=0.05)

    # Перекрёстный эксперимент
    mech = load("mechanism.json")
    for key, vals in [("budget_linear_20", (37.6, 310, 6.4)),
                      ("split_linear_neural_20", (36.2, 137, 1.7)),
                      ("split_neural_linear_15", (36.7, 417, 7.2)),
                      ("budget_neural_15", (35.5, 130, 2.2))]:
        r = mech[key]
        check(f"табл.3 {key} rps", vals[0], r["rps"])
        check(f"табл.3 {key} p95", vals[1], r["light_p95"])
        check(f"табл.3 {key} доля", vals[2], r["light_over_budget_pct"])

    # Цена решения
    dec = load("decision_cost.json")
    check("цена решения: степенной закон", 1.6, dec["power_law"]["microseconds"])
    check("цена решения: сеть", 166, dec["neural"]["microseconds"], tol=0.06)

    # Однородный поток
    io = load("homogeneous_io.json")["levels"]["32"]
    cpu_only = load("homogeneous_cpu.json")["levels"]["32"]
    check("однородный I/O: асинхронный", 231.8, io["all_async"]["rps"])
    check("однородный CPU: синхронный", 55.9, cpu_only["all_sync"]["rps"])
    check("однородный CPU: асинхронный", 8.7, cpu_only["all_async"]["rps"])

    # Доля запросов с ухудшенным откликом
    front = load("budget_frontier.json")
    check("всё в async: доля", 89, front["all_async"]["light_over_budget_pct"])
    check("классификатор: доля (нижняя)", 5,
          front["budget_neural_30"]["light_over_budget_pct"])
    check("классификатор: доля (верхняя)", 7,
          front["budget_power_30"]["light_over_budget_pct"])

    # Объём выборок
    check("обучающая выборка", 384, len(load("cost_samples.json")))
    check("отложенная выборка", 156, len(load("cost_samples_holdout.json")))

    # Числа, которых в тексте быть не должно.
    #
    # Версия статьи с исходной аннотацией — особый случай: числа
    # предварительной серии в ней присутствуют намеренно, потому что
    # аннотацию изменить не удалось, и текст обязан объяснить, откуда
    # расхождение. Остальные устаревшие величины запрещены и там.
    banned_all = ["99,7", "99.7", "6,1 раза", "658", "65 раз", "118 раз", "против 123",
                  "123 мс\u00a0в цикле", "задержка 123",
                  "174 мкс"]
    banned_new = ["39–87", "39-87", "1,8–2,4"]
    for path in sorted(BASE.glob("docs/*.docx")):
        if path.name.startswith("~$"):
            continue
        text = article_text(path)
        legacy_abstract = "v1" in path.name
        for b in banned_all + ([] if legacy_abstract else banned_new):
            if b in text:
                problems.append(f"{path.name}: встречается устаревшее «{b}»")
        if legacy_abstract and "39–87" not in text:
            problems.append(f"{path.name}: исходная аннотация утрачена")

    # Проверка на чужом ПО и производственная трасса
    third = load("thirdparty_postgrest.json")
    check("PostgREST: разброс внутри маршрута", 26.7, third["разброс_раз"])
    check("PostgREST: ошибка по маршруту", 177.1,
          third["ошибка_оценки_по_маршруту_%"])
    check("PostgREST: ошибка по признакам", 33.1,
          third["ошибка_оценки_по_параметру_%"])
    check("PostgREST: один параметр на отложенной",
          56.9, third["на_отложенной_половине"]["один_параметр_ошибка_%"])
    check("PostgREST: все признаки на отложенной",
          38.6, third["на_отложенной_половине"]["все_признаки_ошибка_%"])

    az = load("azure_trace.json")
    check("Azure: функций", 32968, az["функций"])
    check("Azure: доля с разбросом от 10 раз", 44.6,
          az["разброс_p99_к_p50"]["доля_функций_с_разбросом_от_10_раз_%"])
    check("Azure: медиана разброса", 8.0,
          az["разброс_p99_к_p50"]["медиана"])

    # Модуль в nginx и экономика
    mod = load("nginx_module.json")["режимы"]
    check("nginx: медленных у least_conn", 36.9,
          mod["обычная (least_conn)"]["доля_медленнее_200_мс_%"])
    check("nginx: медленных у модуля", 0.0,
          mod["по оценке стоимости"]["доля_медленнее_200_мс_%"])

    cap = load("nginx_capacity.json")["режимы"]
    check("ёмкость least_conn при 100 мс", 37.8,
          cap["обычная (least_conn)"]["ёмкость_при_пороге"]["100.0"]["рпс"])
    check("ёмкость модуля при 100 мс", 59.2,
          cap["по оценке стоимости"]["ёмкость_при_пороге"]["100.0"]["рпс"])

    eco = load("unit_economics.json")["по_порогам"]
    check("экономия в год при 100 мс", 656640,
          eco["100.0"]["разница"]["руб_в_год"])
    check("экономия в год при 200 мс", 0,
          eco["200.0"]["разница"]["руб_в_год"])

    hand = load("handmade.json")
    check("PostgREST вручную: разница", 72, hand["разница_раз"])
    check("PostgREST вручную: самый медленный", 123.1,
          max(z["медиана_мс"] for z in hand["замеры"]))

    # README и заявка — такой же публичный текст, как статьи, и однажды
    # они уже разошлись со статьями после исправления чисел.
    # README и заявка — такой же публичный текст, как статьи, и однажды
    # они уже разошлись со статьями после исправления чисел.
    #
    # PROJECT_STATE намеренно хранит историю проекта, включая отброшенные
    # величины, поэтому в проверку не входит: там устаревшее число — не
    # ошибка, а запись о том, что было исправлено.
    # Публичных текстов стало больше, и однажды README уже разошёлся со
    # статьями после исправления чисел. Проверяются все, кроме
    # PROJECT_STATE: он намеренно хранит историю, включая отброшенные
    # величины, и устаревшее число там — запись, а не ошибка.
    public = ["README.md", "README.en.md", "ПРОВЕРЬ_САМ.md", "loadlens/README.md",
              "loadlens/web/README.md", "docker/README.md",
              "nginx/README.md", "thirdparty/RESULT.md",
              "thirdparty/AZURE.md", "ledentsov/ЗАЯВКА.md",
              "ledentsov/КОНКУРЕНТЫ.md", "ledentsov/ЭКОНОМИКА.md",
              "ledentsov/ВОПРОСЫ_И_ОТВЕТЫ.md"]
    for name in public:
        path = BASE / name
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        for b in ["658 мс", "123 мс", "65 раз", "99,7", "174 мкс", "39–87"]:
            if b in text:
                problems.append(f"{name}: встречается устаревшее «{b}»")

    decks = sorted(BASE.glob("docs/*.pptx")) + \
        sorted(BASE.glob("ledentsov/presentation/*.pptx"))
    for path in decks:
        if path.name.startswith("~$"):
            continue
        text = deck_text(path)
        for b in banned_all + banned_new:
            if b in text:
                problems.append(f"{path.name}: встречается устаревшее «{b}»")

    print(f"Сверено утверждений: {checked}")
    if problems:
        print(f"РАСХОЖДЕНИЙ: {len(problems)}")
        for p in problems:
            print("  •", p)
        return 1
    print("Расхождений нет: каждое число в статьях подтверждается данными.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
