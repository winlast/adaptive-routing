#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Проверка вёрстки презентации без отрисовки.

Открыть файл и посмотреть в этой среде нечем, поэтому вёрстка
проверяется расчётом: ширина строк меряется Liberation Sans,
метрически совместимым с Arial, текст переносится по словам, и
проверяется, помещается ли он в свой блок и не накладываются ли блоки
друг на друга. Проверка приблизительная и заменяет не просмотр файла, а
его отсутствие.
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation

W, H = 13.333, 7.5
FONT_DIR = "/usr/share/fonts/truetype/liberation"
SCALE = 8


def font(size: float, bold: bool):
    name = "LiberationSans-Bold.ttf" if bold else "LiberationSans-Regular.ttf"
    return ImageFont.truetype(f"{FONT_DIR}/{name}", int(size * SCALE))


def text_height(shape, width_in: float) -> float:
    """Высота текста в дюймах при переносе по словам."""
    total = 0.0
    for p in shape.text_frame.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            total += 14 * 1.25 / 72
            continue
        size = next((r.font.size.pt for r in runs if r.font.size), 14)
        bold = bool(runs[0].font.bold)
        f = font(size, bold)
        avail = max(width_in - 0.12, 0.4) * 72
        line, lines = 0.0, 1
        for word in "".join(r.text for r in runs).split():
            w = f.getlength(word + " ") / SCALE
            if line + w > avail and line > 0:
                lines += 1
                line = w
            else:
                line += w
        total += lines * size * 1.25 / 72
    return total


def main(path: Path) -> int:
    prs = Presentation(str(path))
    problems: list[str] = []

    for n, slide in enumerate(prs.slides, 1):
        boxes = []
        for sh in slide.shapes:
            L, T = sh.left / 914400, sh.top / 914400
            Wd, Ht = sh.width / 914400, sh.height / 914400
            txt = sh.text_frame.text.strip() if sh.has_text_frame else ""
            if txt:
                need = text_height(sh, Wd)
                if need > Ht + 0.12:
                    problems.append(
                        f"слайд {n}: «{txt[:34]}…» занимает {need:.2f}\" "
                        f"в блоке {Ht:.2f}\"")
                if T + need > H - 0.05:
                    problems.append(
                        f"слайд {n}: «{txt[:34]}…» выходит за нижний край")
                boxes.append((L, T, Wd, max(Ht, need), txt[:24]))
            elif sh.shape_type == 13:
                boxes.append((L, T, Wd, Ht, "рисунок"))
                if L < -0.05 or T < -0.05 or L + Wd > W + 0.05 or T + Ht > H + 0.05:
                    problems.append(f"слайд {n}: рисунок за краем слайда")

        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                ox = min(a[0] + a[2], b[0] + b[2]) - max(a[0], b[0])
                oy = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
                if ox > 0.2 and oy > 0.2:
                    problems.append(
                        f"слайд {n}: «{a[4]}» и «{b[4]}» накладываются "
                        f"{ox:.2f}×{oy:.2f}\"")

    print(f"Слайдов: {len(prs.slides)}")
    if problems:
        print(f"ЗАМЕЧАНИЯ ({len(problems)}):")
        for p in problems:
            print("  •", p)
        return 1
    print("Вёрстка в порядке: текст помещается, блоки не накладываются.")
    return 0


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else (
        Path(__file__).resolve().parent.parent
        / "docs" / "Презентация Русский инженер Синявский СД.pptx")
    raise SystemExit(main(target))
